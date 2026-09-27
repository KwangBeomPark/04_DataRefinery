"""Focused GUI regression tests for Promotion result-finding workflow.

Tests persistence of export results across tab switches and language changes,
resetting upon choosing a different template, 'Open output folder' button state,
and folder action error handling.
"""

from __future__ import annotations

import tempfile
import tkinter as tk
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl

from src.data_refinery import DataRefineryApp, _UI_TEXT
from src.promotion_normalizer import (
    PROMOTION_MASTER_COLUMNS,
    SUPPORT_RULE_COLUMNS,
    PromotionExportResult,
)


def _create_dummy_template(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws_master = wb.active
    ws_master.title = "Promotion_Master"
    ws_master.append(list(PROMOTION_MASTER_COLUMNS))
    ws_master.append(["PROMO-001", "Promo 1", "Retail", "Discount", ""])

    ws_rules = wb.create_sheet("Support_Rules")
    ws_rules.append(list(SUPPORT_RULE_COLUMNS))
    ws_rules.append(["RULE-001", "PROMO-001", "MOD-1", "2026-01-01", "2026-01-05", "50", "USD"])
    wb.save(path)


class PromotionGuiTestCase(unittest.TestCase):
    def setUp(self):
        self.root = tk.Tk()
        self.root.withdraw()
        self.app = DataRefineryApp(self.root)
        self.temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.dir_path = Path(self.temp_dir.name)

        self.template_a = self.dir_path / "template_a.xlsx"
        self.template_b = self.dir_path / "template_b.xlsx"
        _create_dummy_template(self.template_a)
        _create_dummy_template(self.template_b)

        # Switch to Promotion tab
        self.app._select_task_tab(self.app.promotion_tab)

    def tearDown(self):
        try:
            for after_id in self.root.tk.splitlist(self.root.tk.eval("after info")):
                try:
                    self.root.after_cancel(after_id)
                except Exception:
                    pass
        except Exception:
            pass
        self.root.destroy()

    def _simulate_export_success(self, template_path: Path, daily_rows: int = 5) -> PromotionExportResult:
        self.app.promotion_filepath.set(str(template_path))
        self.app._load_promotion_template()
        daily_path = template_path.parent / f"promotion_daily_support_{template_path.stem}.csv"
        daily_path.touch()
        result = PromotionExportResult(
            master_path=template_path.parent / f"promotion_master_{template_path.stem}.csv",
            rules_path=template_path.parent / f"promotion_support_rules_{template_path.stem}.csv",
            daily_path=daily_path,
            daily_rows=daily_rows,
            overlapping_rule_pairs=0,
        )
        result.master_path.touch()
        result.rules_path.touch()
        self.app._on_promotion_success(result)
        return result


class TestPromotionPersistence(PromotionGuiTestCase):
    def test_export_result_persists_across_tab_switches(self):
        result = self._simulate_export_success(self.template_a, daily_rows=42)

        # Verify summary displayed immediately after export
        initial_summary = self.app.result_text.get("1.0", tk.END)
        self.assertIn(result.daily_path.name, initial_summary)
        self.assertIn(result.master_path.name, initial_summary)
        self.assertIn(result.rules_path.name, initial_summary)
        self.assertIn(str(result.daily_path.parent.resolve()), initial_summary)
        self.assertIn("42", initial_summary)

        # Switch to CSV tab
        self.app._select_task_tab(self.app.csv_tab)
        self.assertEqual(self.app._selected_task_id(), "csv")
        csv_text = self.app.result_text.get("1.0", tk.END)
        self.assertNotIn(result.daily_path.name, csv_text)

        # Switch back to Promotion tab: success summary must persist, NOT input preview
        self.app._select_task_tab(self.app.promotion_tab)
        self.assertEqual(self.app._selected_task_id(), "promotion")
        restored_summary = self.app.result_text.get("1.0", tk.END)

        self.assertIn(result.daily_path.name, restored_summary)
        self.assertIn(result.master_path.name, restored_summary)
        self.assertIn(result.rules_path.name, restored_summary)
        self.assertIn(str(result.daily_path.parent.resolve()), restored_summary)
        self.assertIn("42", restored_summary)
        # Should not revert to preview table header or input issues
        self.assertNotIn("Preview (first", restored_summary)

        # Switch to Aggregator tab and back
        self.app._select_task_tab(self.app.aggregator_tab)
        self.app._select_task_tab(self.app.promotion_tab)
        agg_restored = self.app.result_text.get("1.0", tk.END)
        self.assertIn(result.daily_path.name, agg_restored)

    def test_export_result_persists_and_localizes_across_language_changes(self):
        result = self._simulate_export_success(self.template_a, daily_rows=10)
        output_dir = str(result.daily_path.parent.resolve())

        # English (default)
        self.app.language.set("English")
        self.app._apply_language()
        en_text = self.app.result_text.get("1.0", tk.END)
        self.assertIn("Created 10 daily support rows.", en_text)
        self.assertIn(f"Location: {output_dir}", en_text)
        self.assertIn("Files created", en_text)
        self.assertIn(result.daily_path.name, en_text)
        self.assertEqual(self.app.promotion_open_folder_button.cget("text"), "Open output folder")
        self.assertEqual(str(self.app.promotion_open_folder_button.cget("state")), "normal")

        # Switch to Korean
        self.app.language.set("한국어")
        self.app._apply_language()
        ko_text = self.app.result_text.get("1.0", tk.END)
        self.assertIn("일별 지원금 행 10개를 만들었습니다.", ko_text)
        self.assertIn(f"저장 위치: {output_dir}", ko_text)
        self.assertIn("생성된 파일", ko_text)
        self.assertIn(result.daily_path.name, ko_text)
        self.assertEqual(self.app.promotion_open_folder_button.cget("text"), "결과 폴더 열기")
        self.assertEqual(str(self.app.promotion_open_folder_button.cget("state")), "normal")

        # Switch to Polish
        self.app.language.set("Polski")
        self.app._apply_language()
        pl_text = self.app.result_text.get("1.0", tk.END)
        self.assertIn("Utworzono 10 dziennych wierszy dopłat.", pl_text)
        self.assertIn(f"Lokalizacja: {output_dir}", pl_text)
        self.assertIn("Utworzone pliki", pl_text)
        self.assertIn(result.daily_path.name, pl_text)
        self.assertEqual(self.app.promotion_open_folder_button.cget("text"), "Otwórz folder wynikowy")
        self.assertEqual(str(self.app.promotion_open_folder_button.cget("state")), "normal")


class TestPromotionNewTemplateReset(PromotionGuiTestCase):
    def test_choosing_different_template_resets_export_result_and_disables_button(self):
        result_a = self._simulate_export_success(self.template_a, daily_rows=5)
        self.assertEqual(str(self.app.promotion_open_folder_button.cget("state")), "normal")
        self.assertIsNotNone(self.app._last_promotion_result)

        # Choose a different template via browse
        with patch("src.data_refinery.filedialog.askopenfilename", return_value=str(self.template_b)):
            self.app.browse_promotion_template()

        # The previous export result should be cleared
        self.assertIsNone(self.app._last_promotion_result)
        self.assertEqual(str(self.app.promotion_open_folder_button.cget("state")), "disabled")

        # Result panel should now show the new template preview, not old export summary
        text_after_switch = self.app.result_text.get("1.0", tk.END)
        self.assertNotIn(result_a.daily_path.name, text_after_switch)
        self.assertIn("Preview", text_after_switch)

        # Tab switch must not resurrect old template A result
        self.app._select_task_tab(self.app.csv_tab)
        self.app._select_task_tab(self.app.promotion_tab)
        switched_text = self.app.result_text.get("1.0", tk.END)
        self.assertNotIn(result_a.daily_path.name, switched_text)
        self.assertEqual(str(self.app.promotion_open_folder_button.cget("state")), "disabled")

    def test_filepath_trace_resets_when_different_template_is_set_directly(self):
        self._simulate_export_success(self.template_a, daily_rows=5)
        self.assertEqual(str(self.app.promotion_open_folder_button.cget("state")), "normal")

        # Directly setting a new path via StringVar
        self.app.promotion_filepath.set(str(self.template_b))
        self.assertIsNone(self.app._last_promotion_result)
        self.assertEqual(str(self.app.promotion_open_folder_button.cget("state")), "disabled")

    def test_does_not_open_outdated_result_after_new_template_chosen(self):
        self._simulate_export_success(self.template_a, daily_rows=5)

        # Choose different template
        with patch("src.data_refinery.filedialog.askopenfilename", return_value=str(self.template_b)):
            self.app.browse_promotion_template()

        with patch("src.data_refinery.open_containing_folder") as mock_open:
            opened = self.app.open_promotion_output_folder()
            self.assertFalse(opened)
            mock_open.assert_not_called()


class TestPromotionButtonState(PromotionGuiTestCase):
    def test_button_initially_disabled(self):
        self.assertEqual(str(self.app.promotion_open_folder_button.cget("state")), "disabled")

    def test_button_disabled_when_template_loaded_without_export(self):
        self.app.promotion_filepath.set(str(self.template_a))
        self.app._load_promotion_template()
        self.assertEqual(str(self.app.promotion_open_folder_button.cget("state")), "disabled")

    def test_button_enabled_after_successful_export(self):
        self._simulate_export_success(self.template_a, daily_rows=5)
        self.assertEqual(str(self.app.promotion_open_folder_button.cget("state")), "normal")

    def test_button_disabled_while_processing(self):
        self._simulate_export_success(self.template_a, daily_rows=5)
        self.assertEqual(str(self.app.promotion_open_folder_button.cget("state")), "normal")

        # Disable during processing
        self.app._set_promotion_controls_enabled(False)
        self.assertEqual(str(self.app.promotion_open_folder_button.cget("state")), "disabled")

        # Re-enable after processing finishes
        self.app._set_promotion_controls_enabled(True)
        self.assertEqual(str(self.app.promotion_open_folder_button.cget("state")), "normal")

    def test_button_alias_exists(self):
        self.assertIs(self.app.btn_promo_open_folder, self.app.promotion_open_folder_button)


class TestPromotionFolderAction(PromotionGuiTestCase):
    def test_open_output_folder_selects_daily_support_file(self):
        result = self._simulate_export_success(self.template_a, daily_rows=5)

        with patch("src.data_refinery.open_containing_folder", return_value=True) as mock_open:
            opened = self.app.open_promotion_output_folder()
            self.assertTrue(opened)
            mock_open.assert_called_once_with(str(result.daily_path))

    def test_open_output_folder_handles_failure_gracefully(self):
        result = self._simulate_export_success(self.template_a, daily_rows=5)

        with patch("src.data_refinery.open_containing_folder", return_value=False) as mock_open, \
             patch("src.data_refinery.messagebox.showwarning") as mock_warn:
            opened = self.app.open_promotion_output_folder()
            self.assertFalse(opened)
            mock_open.assert_called_once_with(str(result.daily_path))
            mock_warn.assert_called_once()
            self.assertEqual(mock_warn.call_args[0][0], self.app._ui("promo_msg_open_failed_title"))
            self.assertIn(str(result.daily_path), mock_warn.call_args[0][1])

    def test_open_output_folder_with_missing_file_handles_gracefully(self):
        result = self._simulate_export_success(self.template_a, daily_rows=5)
        # Delete the daily file to simulate missing file
        if result.daily_path.exists():
            result.daily_path.unlink()

        with patch("src.data_refinery.messagebox.showwarning") as mock_warn:
            opened = self.app.open_promotion_output_folder()
            self.assertFalse(opened)
            mock_warn.assert_called_once()

    def test_open_output_folder_before_export_returns_false_silently(self):
        with patch("src.data_refinery.open_containing_folder") as mock_open, \
             patch("src.data_refinery.messagebox.showwarning") as mock_warn:
            opened = self.app.open_promotion_output_folder()
            self.assertFalse(opened)
            mock_open.assert_not_called()
            mock_warn.assert_not_called()


class TestPromotionTranslations(unittest.TestCase):
    def test_all_new_promotion_keys_exist_in_all_languages(self):
        expected_keys = {
            "promo_summary_location",
            "promo_open_folder",
            "promo_msg_open_failed_title",
            "promo_msg_open_failed",
        }
        for lang in ("en", "ko", "pl"):
            for key in expected_keys:
                self.assertIn(
                    key,
                    _UI_TEXT[lang],
                    msg=f"Key {key!r} missing in language {lang!r}",
                )

    def test_placeholders_match_across_languages(self):
        import string

        def placeholders(template):
            return {field for _, field, _, _ in string.Formatter().parse(template) if field}

        for key in ("promo_summary_location", "promo_msg_open_failed"):
            expected = placeholders(_UI_TEXT["en"][key])
            for lang in ("ko", "pl"):
                self.assertEqual(
                    placeholders(_UI_TEXT[lang][key]),
                    expected,
                    msg=f"Placeholder mismatch for {key!r} in {lang!r}",
                )


if __name__ == "__main__":
    unittest.main()
