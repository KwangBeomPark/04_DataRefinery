"""Excel template generator for Data Refinery DWH & Data Mart."""

from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def create_dwh_excel_template(dest_path: Path | str) -> Path:
    """Generate a clean, styled Excel template with Customer, Product, and P&L sheets."""
    dest = Path(dest_path).resolve()
    dest.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()

    # Styling definitions
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10)
    border_side = Side(border_style="thin", color="D1D5DB")
    cell_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

    # 1. Customer Master Sheet
    ws_cust = wb.active
    ws_cust.title = "master_customers"
    cust_headers = ["bill_to_code", "customer_name", "channel_group", "country_code", "city"]
    ws_cust.append(cust_headers)
    ws_cust.append(["000100200301", "MediaExpert Polska", "Retail", "PL", "Warszawa"])
    ws_cust.append(["000100200302", "RTV Euro AGD", "Retail", "PL", "Gdansk"])
    ws_cust.append(["000100200303", "Allegro Direct PL", "Online", "PL", "Poznan"])

    # 2. Product Master Sheet
    ws_prod = wb.create_sheet(title="master_products")
    prod_headers = ["model_code", "product_name", "division", "category"]
    ws_prod.append(prod_headers)
    ws_prod.append(["OLED65C6P", "OLED TV 65 C6", "HE", "OLED TV"])
    ws_prod.append(["QNED55B6P", "QNED TV 55 B6", "HE", "QNED TV"])
    ws_prod.append(["WM9000PL", "Eco Washer 9kg", "HA", "Washing Machine"])

    # 3. Raw P&L Sheet
    ws_pnl = wb.create_sheet(title="raw_pnl")
    pnl_headers = [
        "period", "entity_code", "bill_to_code", "model_code", "currency_code", "sales_qty",
        "REV_001", "DED_001", "COG_001", "LOG_001", "PRM_001", "WAR_001", "FIX_001"
    ]
    ws_pnl.append(pnl_headers)
    ws_pnl.append([
        "2026-01", "PL", "000100200301", "OLED65C6P", "PLN", 150,
        150000.0, 3000.0, 75000.0, 4500.0, 6000.0, 1500.0, 12000.0
    ])
    ws_pnl.append([
        "2026-01", "PL", "000100200302", "QNED55B6P", "PLN", 80,
        80000.0, 1600.0, 40000.0, 2400.0, 3200.0, 800.0, 6400.0
    ])

    for ws in [ws_cust, ws_prod, ws_pnl]:
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = cell_border
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = data_font
                cell.border = cell_border
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    wb.save(str(dest))
    return dest
