"""CSV and Excel cleanup service with no Tkinter dependency."""

from __future__ import annotations

import csv
import codecs
import decimal
import math
import os
import re
import uuid
from collections import Counter, deque
from dataclasses import dataclass
from datetime import datetime
from itertools import islice
from typing import Callable, Iterable, Iterator


_PANDAS = None
_ENCODING_SAMPLE_BYTES = 1024 * 1024
_PROGRESS_ROW_INTERVAL = 5000
_EXCEL_MAX_DATA_ROWS = 1_048_575
_DATE_HINT_RE = re.compile(r"\d{1,4}[-/.]\d{1,2}[-/.]\d{1,4}")
_LINEBREAK_RE = re.compile("[\r\n\x0b\x0c\x85\u2028\u2029]+")


class ParsedNumber:
    __slots__ = ["value", "orig_decimals", "orig_text"]

    def __init__(self, value: decimal.Decimal, orig_decimals: int, orig_text: str):
        self.value = value
        self.orig_decimals = orig_decimals
        self.orig_text = orig_text


class CsvNoTableError(ValueError):
    """The source has no row that can serve as a table header."""


class CsvNoDataError(ValueError):
    """The source has a header but no data rows."""


@dataclass(frozen=True)
class CsvProcessingOptions:
    file_path: str
    delimiter: str
    number_mode: str
    output_format: str
    max_columns: int


@dataclass(frozen=True)
class CsvProcessingResult:
    out_path: str
    rows: int
    columns: int
    garbage_skipped: int
    numbers: int
    dates: int
    flattened: int
    repaired: int
    large_numbers_as_text: int
    encoding: str


def _get_pandas():
    """Import pandas only when a selected file actually needs processing."""
    global _PANDAS
    if _PANDAS is None:
        import pandas as pandas_module

        _PANDAS = pandas_module
    return _PANDAS


def is_excel(file_path: str) -> bool:
    return os.path.splitext(file_path)[1].lower() in (".xlsx", ".xlsm", ".xls")


def normalize_delimiter(delimiter: str) -> str:
    """Accept common typed spellings for a tab separator."""
    return {"\\t": "\t", "tab": "\t", "TAB": "\t", "Tab": "\t"}.get(delimiter, delimiter)


def detect_encoding(file_path: str) -> str:
    """Pick the most plausible text encoding without silently corrupting text."""
    with open(file_path, "rb") as file:
        raw = file.read(_ENCODING_SAMPLE_BYTES)

    if raw.startswith(b"\xff\xfe") or raw.startswith(b"\xfe\xff"):
        return "utf-16"
    if raw and raw.count(b"\x00") > len(raw) // 4:
        for encoding in ("utf-16-le", "utf-16-be"):
            try:
                raw.decode(encoding)
                return encoding
            except UnicodeDecodeError:
                pass

    def can_decode_prefix(encoding: str) -> bool:
        try:
            codecs.getincrementaldecoder(encoding)().decode(raw, final=False)
            return True
        except UnicodeDecodeError:
            return False

    for encoding in ("utf-8-sig", "utf-8"):
        if can_decode_prefix(encoding):
            return encoding

    best: tuple[str, int] | None = None
    for encoding in ("cp949", "cp1252"):
        try:
            text = codecs.getincrementaldecoder(encoding)().decode(raw, final=False)
        except UnicodeDecodeError:
            continue
        suspicious_count = sum(1 for character in text if ("\x80" <= character <= "\x9f") or character == "�")
        if best is None or suspicious_count < best[1]:
            best = (encoding, suspicious_count)
    return best[0] if best is not None else "latin-1"


def detect_delimiter(file_path: str, sample_rows: int = 50) -> str | None:
    """Infer a delimiter from stable logical CSV record widths."""
    if is_excel(file_path):
        return None
    try:
        encoding = detect_encoding(file_path)
        best_delimiter = None
        best_score = None
        for candidate in (";", ",", "\t", "|"):
            with open(file_path, "r", encoding=encoding, newline="") as file:
                reader = csv.reader(file, delimiter=candidate)
                rows = []
                for row in reader:
                    if row:
                        rows.append(row)
                    if len(rows) >= sample_rows:
                        break
            if not rows:
                continue

            width, count = max(Counter(len(row) for row in rows).items(), key=lambda item: (item[1], item[0]))
            if width <= 1:
                continue
            score = (count / len(rows), count, width)
            if best_score is None or score > best_score:
                best_score = score
                best_delimiter = candidate
        return best_delimiter
    except Exception:
        return None


def read_file_rows(file_path: str, delimiter: str, max_rows: int | None = None) -> tuple[list[list], str]:
    if is_excel(file_path):
        return read_excel_rows(file_path, max_rows), "Excel"
    delimiter = normalize_delimiter(delimiter)
    if not delimiter or len(delimiter) != 1:
        raise ValueError("Delimiter must be a single character.")
    encoding = detect_encoding(file_path)
    with open(file_path, "r", encoding=encoding, newline="") as file:
        reader = csv.reader(file, delimiter=delimiter)
        rows = []
        for index, row in enumerate(reader):
            if max_rows is not None and index >= max_rows:
                break
            rows.append(row)
    return rows, encoding


def read_excel_rows(file_path: str, max_rows: int | None = None) -> list[list]:
    """Read native Excel values while applying the CSV row-shape rules."""
    pandas = _get_pandas()
    frame = pandas.read_excel(file_path, header=None, dtype=object, nrows=max_rows)
    rows = []
    for record in frame.itertuples(index=False, name=None):
        row = []
        for value in record:
            if value is None or (not isinstance(value, str) and pandas.isna(value)):
                row.append("")
            elif isinstance(value, datetime):
                row.append(value.date() if (value.hour, value.minute, value.second) == (0, 0, 0) else value)
            else:
                row.append(value)
        while row and row[-1] == "":
            row.pop()
        rows.append(row)
    return rows


@dataclass
class RepairStatistics:
    repaired: int = 0


def iter_repaired_rows(
    rows: Iterable[list],
    max_columns: int,
    statistics: RepairStatistics | None = None,
) -> Iterator[list]:
    """Rejoin split records while retaining at most 21 source rows in memory."""
    statistics = statistics or RepairStatistics()
    source = iter(rows)
    pushed_back: deque[list] = deque()
    pending: list | None = None

    def next_row() -> list:
        return pushed_back.popleft() if pushed_back else next(source)

    while True:
        try:
            row = next_row()
        except StopIteration:
            break

        if 0 < len(row) < max_columns:
            if len(row) == 1 and pending is not None and len(pending) == max_columns:
                pending[-1] = f"{pending[-1]} {row[0]}".strip()
                statistics.repaired += 1
                continue

            merged = list(row)
            consumed: list[list] = []
            while len(merged) < max_columns and len(consumed) < 20:
                try:
                    following = next_row()
                except StopIteration:
                    break
                if (
                    len(following) >= max_columns
                    or len(merged) + max(len(following), 1) - 1 > max_columns
                ):
                    pushed_back.appendleft(following)
                    break
                consumed.append(following)
                if following:
                    merged[-1] = f"{merged[-1]} {following[0]}".strip()
                    merged.extend(following[1:])

            if len(merged) == max_columns and consumed:
                row = merged
                statistics.repaired += 1
            else:
                pushed_back.extendleft(reversed(consumed))

        if pending is not None:
            yield pending
        pending = row

    if pending is not None:
        yield pending


def repair_split_rows(rows: list[list], max_columns: int) -> tuple[list[list], int]:
    """Compatibility helper for callers that already hold a small row list."""
    statistics = RepairStatistics()
    repaired_rows = list(iter_repaired_rows(rows, max_columns, statistics))
    return repaired_rows, statistics.repaired


def parse_number(value, mode: str):
    """Parse validated English or Polish values while retaining decimal scale."""
    if not isinstance(value, str):
        return value
    text = value.strip()
    if not text:
        return value

    normalized = text.replace("\u00a0", " ").replace("\u202f", " ").replace("\u2212", "-")
    sign = ""
    if normalized.startswith(("-", "+")):
        sign, normalized = normalized[0], normalized[1:]
    if not normalized:
        return value

    if mode == "Polish":
        if not re.fullmatch(r"[0-9 \\.,]+", normalized):
            return value
        integer_part, *fraction_parts = normalized.split(",")
        if len(fraction_parts) > 1:
            return value
        fraction_part = fraction_parts[0] if fraction_parts else None
        if fraction_part == "":
            return value
        if not integer_part and fraction_part:
            integer_part = "0"
        if " " in integer_part and "." in integer_part:
            return value
        separator = " " if " " in integer_part else "." if "." in integer_part else None
        if separator:
            groups = integer_part.split(separator)
            if not groups[0] or len(groups[0]) > 3 or any(len(group) != 3 for group in groups[1:]):
                return value
            integer_part = "".join(groups)
        if fraction_part and ("." in fraction_part or " " in fraction_part):
            return value
    elif mode == "English":
        if not re.fullmatch(r"[0-9\\.,]+", normalized):
            return value
        integer_part, *fraction_parts = normalized.split(".")
        if len(fraction_parts) > 1:
            return value
        fraction_part = fraction_parts[0] if fraction_parts else None
        if fraction_part == "":
            return value
        if not integer_part and fraction_part:
            integer_part = "0"
        if "," in integer_part:
            groups = integer_part.split(",")
            if not groups[0] or len(groups[0]) > 3 or any(len(group) != 3 for group in groups[1:]):
                return value
            integer_part = "".join(groups)
        if fraction_part and "," in fraction_part:
            return value
    else:
        return value

    if len(integer_part) > 1 and integer_part.startswith("0"):
        return value
    decimal_places = len(fraction_part) if fraction_part else 0
    decimal_text = f"{sign}{integer_part}.{fraction_part}" if fraction_part else f"{sign}{integer_part}"
    try:
        return ParsedNumber(decimal.Decimal(decimal_text), decimal_places, value)
    except decimal.InvalidOperation:
        return value


def parse_date(value):
    """Parse only dates whose numeric order is unambiguous."""
    if not isinstance(value, str):
        return value
    text = value.strip()
    if not text or not _DATE_HINT_RE.fullmatch(text):
        return value

    ambiguous_numeric_date = re.fullmatch(r"(\d{1,2})([/-])(\d{1,2})\2(\d{4})", text)
    if ambiguous_numeric_date:
        first = int(ambiguous_numeric_date.group(1))
        second = int(ambiguous_numeric_date.group(3))
        if first <= 12 and second <= 12:
            return value

    for date_format in (
        "%Y-%m-%d", "%Y/%m/%d", "%d.%m.%Y", "%d-%m-%Y", "%d/%m/%Y", "%m-%d-%Y", "%m/%d/%Y"
    ):
        try:
            return datetime.strptime(text, date_format).date()
        except ValueError:
            continue
    return value


def decimal_significant_digits(value: decimal.Decimal) -> int:
    normalized = value.normalize()
    digits = normalized.as_tuple().digits
    return 1 if all(digit == 0 for digit in digits) else len(digits)


def requires_excel_text(value: ParsedNumber) -> bool:
    return decimal_significant_digits(value.value) > 15


def format_csv_value(value, decimal_separator: str):
    if isinstance(value, bool):
        return value
    if isinstance(value, ParsedNumber):
        text = f"{value.value:.{value.orig_decimals}f}"
        return text.replace(".", decimal_separator) if decimal_separator != "." else text
    if isinstance(value, float):
        if math.isnan(value):
            return ""
        text = repr(value)
        return text.replace(".", decimal_separator) if decimal_separator != "." else text
    return value


def write_excel_output(frame, out_path: str) -> int:
    """Write converted values and preserve decimal formatting where Excel can."""
    import openpyxl

    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.append(list(frame.columns))
    large_numbers_as_text = 0
    for row in frame.itertuples(index=False):
        output_row = []
        for value in row:
            if isinstance(value, ParsedNumber):
                if requires_excel_text(value):
                    large_numbers_as_text += 1
                    output_row.append(value.orig_text)
                else:
                    output_row.append(float(value.value))
            else:
                output_row.append(value)
        worksheet.append(output_row)

    for row_index, row in enumerate(frame.itertuples(index=False), start=2):
        for column_index, value in enumerate(row, start=1):
            if isinstance(value, ParsedNumber) and not requires_excel_text(value):
                cell = worksheet.cell(row=row_index, column=column_index)
                cell.number_format = "0" if value.orig_decimals == 0 else "0." + ("0" * value.orig_decimals)
    workbook.save(out_path)
    return large_numbers_as_text


def output_extension(output_format: str) -> str:
    return ".xlsx" if output_format == "Excel (.xlsx)" else ".csv"


def build_output_path(file_path: str, output_format: str, timestamp: datetime | None = None) -> str:
    stamp = (timestamp or datetime.now()).strftime("%Y%m%d_%H%M")
    directory = os.path.dirname(file_path)
    base_name = f"processed_output_{stamp}"
    candidate = os.path.join(directory, f"{base_name}{output_extension(output_format)}")
    sequence = 2
    while os.path.exists(candidate):
        candidate = os.path.join(directory, f"{base_name}_{sequence:02d}{output_extension(output_format)}")
        sequence += 1
    return candidate


def _normalise_row(row: list, max_columns: int) -> list:
    normalised = list(row)
    if len(normalised) < max_columns:
        normalised.extend([""] * (max_columns - len(normalised)))
    return normalised[:max_columns]


def _column_names(header_row: list) -> list[str]:
    names = []
    occurrences: dict[str, int] = {}
    for index, value in enumerate(header_row):
        name = _LINEBREAK_RE.sub(" ", "" if value is None else str(value)).strip() or f"Column{index + 1}"
        if name in occurrences:
            occurrences[name] += 1
            name = f"{name}_{occurrences[name]}"
        else:
            occurrences[name] = 0
        names.append(name)
    return names


@dataclass
class ProcessingStatistics:
    numbers: int = 0
    dates: int = 0
    flattened: int = 0
    large_numbers_as_text: int = 0


def _convert_value(value, number_mode: str, statistics: ProcessingStatistics):
    if not isinstance(value, str):
        return value
    if _LINEBREAK_RE.search(value):
        value = _LINEBREAK_RE.sub(" ", value)
        statistics.flattened += 1
    parsed_date = parse_date(value)
    if not isinstance(parsed_date, str):
        statistics.dates += 1
        return parsed_date
    parsed_number = parse_number(value, number_mode)
    if isinstance(parsed_number, ParsedNumber):
        statistics.numbers += 1
    return parsed_number


def _temporary_output_path(out_path: str) -> str:
    directory = os.path.dirname(out_path)
    name = os.path.basename(out_path)
    return os.path.join(directory, f".{name}.{uuid.uuid4().hex}.tmp")


def _find_text_header_index(
    file_path: str,
    delimiter: str,
    encoding: str,
    max_columns: int,
) -> int:
    with open(file_path, "r", encoding=encoding, newline="") as source:
        for index, row in enumerate(csv.reader(source, delimiter=delimiter)):
            if len(row) == max_columns:
                return index
    return 0


def _excel_stream_row(values: list, worksheet, statistics: ProcessingStatistics) -> list:
    from openpyxl.cell import WriteOnlyCell

    output_row = []
    for value in values:
        if not isinstance(value, ParsedNumber):
            output_row.append(value)
            continue
        if requires_excel_text(value):
            statistics.large_numbers_as_text += 1
            output_row.append(value.orig_text)
            continue
        cell = WriteOnlyCell(worksheet, value=float(value.value))
        cell.number_format = "0" if value.orig_decimals == 0 else "0." + ("0" * value.orig_decimals)
        output_row.append(cell)
    return output_row


def _process_text_stream(
    options: CsvProcessingOptions,
    report: Callable[[int, str], None],
) -> CsvProcessingResult:
    """Process delimited text with bounded memory and atomic final publication."""
    delimiter = normalize_delimiter(options.delimiter)
    if not delimiter or len(delimiter) != 1:
        raise ValueError("Delimiter must be a single character.")

    encoding = detect_encoding(options.file_path)
    header_index = _find_text_header_index(
        options.file_path,
        delimiter,
        encoding,
        options.max_columns,
    )
    report(5, f"scanning:{header_index}")

    out_path = build_output_path(options.file_path, options.output_format)
    temporary_path = _temporary_output_path(out_path)
    source_size = max(1, os.path.getsize(options.file_path))
    processing = ProcessingStatistics()
    repair = RepairStatistics()
    body_row_count = 0
    output_file = None
    workbook = None

    try:
        with open(options.file_path, "r", encoding=encoding, newline="") as source:
            source_rows = islice(csv.reader(source, delimiter=delimiter), header_index, None)
            repaired_rows = iter_repaired_rows(source_rows, options.max_columns, repair)
            try:
                header_row = next(repaired_rows)
            except StopIteration as error:
                raise CsvNoTableError() from error

            column_names = _column_names(_normalise_row(header_row, options.max_columns))
            if options.output_format == "Excel (.xlsx)":
                import openpyxl

                workbook = openpyxl.Workbook(write_only=True)
                worksheet = workbook.create_sheet("Cleaned_Data")
                worksheet.append(column_names)

                def write_row(values: list) -> None:
                    worksheet.append(_excel_stream_row(values, worksheet, processing))
            else:
                output_file = open(temporary_path, "w", encoding="utf-8-sig", newline="")
                writer = csv.writer(
                    output_file,
                    delimiter=";" if options.number_mode == "Polish" else ",",
                )
                writer.writerow(column_names)
                decimal_separator = "," if options.number_mode == "Polish" else "."

                def write_row(values: list) -> None:
                    writer.writerow(
                        format_csv_value(value, decimal_separator)
                        for value in values
                    )

            for row in repaired_rows:
                if (
                    options.output_format == "Excel (.xlsx)"
                    and body_row_count >= _EXCEL_MAX_DATA_ROWS
                ):
                    raise ValueError(
                        "The result exceeds Excel's row limit. Select CSV output."
                    )
                normalised = _normalise_row(row, options.max_columns)
                converted = [
                    _convert_value(value, options.number_mode, processing)
                    for value in normalised
                ]
                write_row(converted)
                body_row_count += 1
                if body_row_count % _PROGRESS_ROW_INTERVAL == 0:
                    position = source.buffer.tell()
                    percent = 10 + int(70 * min(1.0, position / source_size))
                    report(percent, f"scanning:{body_row_count}")

        if body_row_count == 0:
            raise CsvNoDataError()

        report(85, "saving")
        if output_file is not None:
            output_file.close()
            output_file = None
        else:
            workbook.save(temporary_path)
            workbook.close()
            workbook = None
        os.replace(temporary_path, out_path)
    except Exception:
        if output_file is not None:
            output_file.close()
        if workbook is not None:
            workbook.close()
        try:
            os.unlink(temporary_path)
        except FileNotFoundError:
            pass
        raise

    return CsvProcessingResult(
        out_path=out_path,
        rows=body_row_count,
        columns=options.max_columns,
        garbage_skipped=header_index,
        numbers=processing.numbers,
        dates=processing.dates,
        flattened=processing.flattened,
        repaired=repair.repaired,
        large_numbers_as_text=processing.large_numbers_as_text,
        encoding=encoding,
    )


def _process_excel_in_memory(
    options: CsvProcessingOptions,
    report: Callable[[int, str], None],
) -> CsvProcessingResult:
    """Keep native Excel processing compatible; large delimited text uses streaming."""
    rows, encoding = read_file_rows(options.file_path, options.delimiter)
    report(10, f"scanning:{len(rows)}")

    start_index = next((index for index, row in enumerate(rows) if len(row) == options.max_columns), 0)
    garbage_skipped = start_index
    if start_index:
        del rows[:start_index]
    if not rows:
        raise CsvNoTableError()

    header_row = _normalise_row(rows[0], options.max_columns)
    body_row_count = len(rows) - 1
    if body_row_count <= 0:
        raise CsvNoDataError()
    column_names = _column_names(header_row)
    pandas = _get_pandas()
    frame = pandas.DataFrame(
        (_normalise_row(row, options.max_columns) for row in rows[1:]),
        columns=column_names,
    )
    del rows

    statistics = ProcessingStatistics()
    column_count = len(frame.columns)
    for column_index, column in enumerate(frame.columns):
        frame[column] = frame[column].map(
            lambda value: _convert_value(value, options.number_mode, statistics)
        )
        report(
            10 + int(70 * (column_index + 1) / column_count),
            f"converting:{column_index + 1}:{column_count}",
        )

    frame.fillna("", inplace=True)
    report(85, "saving")
    out_path = build_output_path(options.file_path, options.output_format)
    temporary_path = _temporary_output_path(out_path)
    try:
        if options.output_format == "Excel (.xlsx)":
            statistics.large_numbers_as_text = write_excel_output(frame, temporary_path)
        else:
            decimal_separator = "," if options.number_mode == "Polish" else "."
            for column in frame.columns:
                frame[column] = frame[column].map(
                    lambda value: format_csv_value(value, decimal_separator)
                )
            frame.to_csv(
                temporary_path,
                sep=";" if options.number_mode == "Polish" else ",",
                index=False,
                encoding="utf-8-sig",
            )
        os.replace(temporary_path, out_path)
    except Exception:
        try:
            os.unlink(temporary_path)
        except FileNotFoundError:
            pass
        raise

    return CsvProcessingResult(
        out_path=out_path,
        rows=body_row_count,
        columns=column_count,
        garbage_skipped=garbage_skipped,
        numbers=statistics.numbers,
        dates=statistics.dates,
        flattened=statistics.flattened,
        repaired=0,
        large_numbers_as_text=statistics.large_numbers_as_text,
        encoding=encoding,
    )


def process_csv_file(
    options: CsvProcessingOptions,
    progress: Callable[[int, str], None] | None = None,
) -> CsvProcessingResult:
    """Repair, normalize, and export one source file without touching UI state."""
    def report(percent: int, event: str) -> None:
        if progress is not None:
            progress(percent, event)

    report(0, "reading")
    if is_excel(options.file_path):
        return _process_excel_in_memory(options, report)
    return _process_text_stream(options, report)
